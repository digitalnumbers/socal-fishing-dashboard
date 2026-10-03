#!/usr/bin/env python3
"""Fail-closed validation gate for a staged SoCal Fishing Dashboard build."""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd


REQUIRED = {
    "buoy_daily.csv": (["station", "date"], 1000),
    "mur_history.csv": (["zone_id", "date"], 1000),
    "shore_sst_daily.csv": (["tide_station", "date"], 1000),
    "conditions_daily.csv": (["zone_id", "date"], 100),
    "scores_daily.csv": (["date", "zone_id", "species_id"], 1000),
    "tide_daily.csv": (["tide_station", "date"], 30),
    "extended_outlook.csv": (["date", "zone_id"], 150),
    "extended_scores.csv": (["date", "zone_id", "species_id"], 1000),
}
HISTORY = {"buoy_daily.csv": .995, "mur_history.csv": .995, "shore_sst_daily.csv": .995,
           "catch_reports.csv": .95}


def fail(errors: list[str], message: str) -> None:
    errors.append(message)


def load_data_js(path: Path) -> dict:
    raw = path.read_text(encoding="utf-8")
    match = re.fullmatch(r"\s*window\.DATA\s*=\s*(\{.*\})\s*;\s*", raw, re.S)
    if not match:
        raise ValueError("data.js is not a single valid window.DATA assignment")
    return json.loads(match.group(1))


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def validate(root: Path, baseline: Path | None, today: date) -> dict:
    csv_dir, app = root / "dataset" / "csv", root / "app"
    errors, warnings, counts = [], [], {}
    frames: dict[str, pd.DataFrame] = {}
    for name, (key, floor) in REQUIRED.items():
        path = csv_dir / name
        if not path.exists() or path.stat().st_size == 0:
            fail(errors, f"missing or empty required output: {name}")
            continue
        try:
            df = pd.read_csv(path)
            frames[name] = df
            counts[name] = len(df)
            if len(df) < floor:
                fail(errors, f"{name} has {len(df)} rows; minimum is {floor}")
            absent = [c for c in key if c not in df]
            if absent:
                fail(errors, f"{name} missing primary-key columns: {absent}")
            elif df.duplicated(key).any():
                fail(errors, f"{name} contains duplicate primary keys {key}")
        except Exception as exc:
            fail(errors, f"{name} is unparsable: {exc}")

    if baseline and baseline.exists():
        for name, ratio in HISTORY.items():
            current, old = csv_dir / name, baseline / name
            if current.exists() and old.exists():
                try:
                    n, n0 = sum(1 for _ in open(current, encoding="utf-8")) - 1, sum(1 for _ in open(old, encoding="utf-8")) - 1
                    if n < n0 * ratio:
                        fail(errors, f"{name} regressed from {n0} to {n} rows")
                    new_dates = pd.to_datetime(
                        pd.read_csv(current, usecols=["date"])["date"], errors="coerce"
                    ).dropna()
                    old_dates = pd.to_datetime(
                        pd.read_csv(old, usecols=["date"])["date"], errors="coerce"
                    ).dropna()
                    if not new_dates.empty and not old_dates.empty:
                        if new_dates.min() > old_dates.min():
                            fail(
                                errors,
                                f"{name} historical start regressed from "
                                f"{old_dates.min().date()} to {new_dates.min().date()}",
                            )
                        if new_dates.max() < old_dates.max():
                            fail(
                                errors,
                                f"{name} latest date regressed from "
                                f"{old_dates.max().date()} to {new_dates.max().date()}",
                            )
                except UnicodeDecodeError:
                    fail(errors, f"{name} could not be checked for historical regression")
                except (KeyError, ValueError) as exc:
                    fail(errors, f"{name} date coverage could not be checked: {exc}")

    scores = frames.get("scores_daily.csv")
    if scores is not None:
        for col in ("score", "seasonal_norm_score", "annual_percentile"):
            if col in scores and ((scores[col].dropna() < -0.01) | (scores[col].dropna() > 100.01)).any():
                fail(errors, f"{col} is outside 0-100 tolerance")
    ext_scores = frames.get("extended_scores.csv")
    if ext_scores is not None:
        for col in ("score_outlook", "score_lo", "score_hi", "climatological_score"):
            if col in ext_scores and ((ext_scores[col].dropna() < -0.01) | (ext_scores[col].dropna() > 100.01)).any():
                fail(errors, f"{col} is outside 0-100 tolerance")

    cond = frames.get("conditions_daily.csv")
    if cond is not None:
        dates = pd.to_datetime(cond.date).dt.date
        if cond.zone_id.nunique() != 9:
            fail(errors, f"conditions_daily has {cond.zone_id.nunique()} zones, expected 9")
        if dates.max() < today + timedelta(days=6):
            fail(errors, f"near-term forecast ends {dates.max()}, before {today + timedelta(days=6)}")
        if dates.min() > today - timedelta(days=13):
            fail(errors, "near-term hindcast coverage is shorter than 14 days")
    ext = frames.get("extended_outlook.csv")
    if ext is not None:
        ext_dates = pd.to_datetime(ext.date).dt.date
        if ext_dates.max() < today + timedelta(days=30):
            fail(errors, f"extended outlook ends {ext_dates.max()}, before {today + timedelta(days=30)}")
        leads = pd.to_numeric(ext.get("lead_days"), errors="coerce")
        if (leads < 7).any():
            fail(errors, "extended_outlook overlaps the lead 0-6 seven-day view")
        if "confidence_tier" in ext:
            tiers = pd.to_numeric(ext["confidence_tier"], errors="coerce")
            if (tiers[(leads >= 7) & (leads <= 14)] != 2).any():
                fail(errors, "lead days 7-14 must use moderate-confidence tier 2")
            if (tiers[leads >= 15] != 3).any():
                fail(errors, "lead days 15-30 must use outlook tier 3")
        far = ext[pd.to_numeric(ext.get("lead_days"), errors="coerce") >= 15]
        # Raw model columns may contain Open-Meteo's last valid endpoint at lead 15.
        # Publication is governed by the projected fields and sea_state_published.
        for col in ("wave_proj_ft", "wind_proj_kt"):
            if col in far and far[col].notna().any():
                fail(errors, f"{col} publishes deterministic values at lead >=15")
        if "sea_state_published" in far and far.sea_state_published.astype(str).str.lower().isin(["true", "1"]).any():
            fail(errors, "sea_state_published is true at lead >=15")

    status_path = csv_dir / "source_status.json"
    try:
        status = json.loads(status_path.read_text())
        if not isinstance(status.get("sources"), list) or not status["sources"]:
            fail(errors, "source_status.json has no source rows")
        for row in status.get("sources", []):
            if (row.get("freshness") == "failed"
                    and row.get("deployment_critical")
                    and not row.get("used_cached_data")):
                fail(errors, f"failed source has no last-known-good fallback: {row.get('source_identifier')}")
    except Exception as exc:
        fail(errors, f"source_status.json is missing or invalid: {exc}")

    try:
        payload = load_data_js(app / "data.js")
        for key in ("conditions", "scores", "extended", "extended_scores", "source_status"):
            if not payload.get(key):
                fail(errors, f"app/data.js missing or empty key: {key}")
    except Exception as exc:
        fail(errors, f"app/data.js invalid: {exc}")

    for rel in ("index.html", "app.js", "data.js", "favicon.svg"):
        if not (app / rel).exists():
            fail(errors, f"site asset missing: app/{rel}")
    for source in csv_dir.glob("*.csv"):
        name = source.name
        staged = app / "downloads" / "csv" / name
        if not staged.exists() or sha(staged) != sha(source):
            fail(errors, f"download copy missing or inconsistent: {name}")
    workbook = root / "dataset" / "socal_fishing_dataset.xlsx"
    try:
        wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
        expected = {"conditions_daily", "scores_daily", "extended_outlook", "extended_scores"}
        if not expected.issubset(set(wb.sheetnames)):
            fail(errors, f"workbook missing sheets: {sorted(expected - set(wb.sheetnames))}")
        for sheet in expected & set(wb.sheetnames):
            csv_name = f"{sheet}.csv"
            frame = frames.get(csv_name)
            if frame is not None:
                workbook_rows = max(0, wb[sheet].max_row - 1)
                if workbook_rows != len(frame):
                    fail(
                        errors,
                        f"workbook sheet {sheet} has {workbook_rows} data rows; "
                        f"{csv_name} has {len(frame)}",
                    )
        wb.close()
    except Exception as exc:
        fail(errors, f"workbook invalid: {exc}")
    if not (app / "downloads" / "socal_fishing_dataset.xlsx").exists() or (
        workbook.exists() and sha(workbook) != sha(app / "downloads" / "socal_fishing_dataset.xlsx")
    ):
        fail(errors, "download workbook is missing or inconsistent")

    result = {"ok": not errors, "validated_for_local_date": today.isoformat(),
              "errors": errors, "warnings": warnings, "row_counts": counts}
    print(json.dumps(result, indent=2))
    return result


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    ap.add_argument("--baseline", type=Path)
    ap.add_argument("--today", type=date.fromisoformat, default=date.today())
    ap.add_argument("--report", type=Path)
    args = ap.parse_args()
    result = validate(args.root.resolve(), args.baseline.resolve() if args.baseline else None, args.today)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(result, indent=2) + "\n")
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
